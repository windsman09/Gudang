
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_excel(headers: list[str], rows: list[tuple], sheet_name: str) -> BytesIO:
    """Membuat file Excel di memori dan mengembalikannya sebagai BytesIO."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    worksheet.append(headers)

    for row in rows:
        worksheet.append(list(row))

    # Format header
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB",
        )
        cell.alignment = Alignment(horizontal="center")

    # Lebar kolom otomatis berdasarkan isi
    for column in worksheet.columns:
        max_length = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column
        )
        worksheet.column_dimensions[
            get_column_letter(column[0].column)
        ].width = min(max_length + 3, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

